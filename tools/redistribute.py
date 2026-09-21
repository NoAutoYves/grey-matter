"""
Redistribute multiple-choice options inside .docx exercise files so that
correct answers hit a target A/B/C/D distribution. Files are rewritten in place.

Only files whose name contains " - Exercises.docx" are processed.

If any question cannot be parsed, its location is listed at the top of the
document in a "Questions not parsed" section, and the file is still saved
with the questions that did parse.
"""

import sys
import random
from pathlib import Path
from collections import deque
from docx import Document


BASE_DIR = Path(r"C:\Users\madon\Documents\GM\MATH\GRADE 10")
# BASE_DIR = Path(r"C:\Users\madon\Documents\GM\MATH\GRADE 12\test")
TARGETS = [13, 12, 13, 12]  # A, B, C, D
NAME_FILTER = " - Exercises.docx"


def find_questions(paragraphs):
    """
    Walk the paragraph list. A question starts at a paragraph whose text is
    'Question N:' and ends at the paragraph beginning with 'Answer:'.

    Returns a list of dicts describing each parsed question.
    Also returns a list of 'Question N:' labels for any block that could not
    be parsed.
    """
    parsed = []
    unparsed = []
    i = 0
    n = len(paragraphs)

    while i < n:
        text = paragraphs[i].text.strip()
        if not text.startswith("Question "):
            i += 1
            continue

        start = i
        label = text

        # Advance to find A), B), C), D) and Answer: in order.
        body_line = None
        a_i = b_i = c_i = d_i = ans_i = None
        j = i + 1
        while j < n:
            t = paragraphs[j].text.strip()

            if t.startswith("A)") and a_i is None:
                a_i = j
            elif t.startswith("B)") and a_i is not None and b_i is None:
                b_i = j
            elif t.startswith("C)") and b_i is not None and c_i is None:
                c_i = j
            elif t.startswith("D)") and c_i is not None and d_i is None:
                d_i = j
            elif t.startswith("Answer:") and d_i is not None and ans_i is None:
                ans_i = j
                break
            elif t.startswith("Question ") and a_i is None:
                # We hit another question before this one had options —
                # malformed, give up on this block.
                break

            j += 1

        if a_i is None or b_i is None or c_i is None or d_i is None or ans_i is None:
            unparsed.append(label)
            i = j + 1 if j < n else n
            continue

        # The question body is whatever non-blank text sits between the
        # 'Question N:' line and the 'A)' line. Usually just one line.
        body_text = paragraphs[a_i - 1].text if a_i - 1 > start else ""

        answer_text = paragraphs[ans_i].text.strip()
        answer = answer_text.split(":", 1)[1].strip().upper()
        if answer not in ("A", "B", "C", "D"):
            unparsed.append(label)
            i = ans_i + 1
            continue

        def strip_prefix(s, prefix):
            s = s.strip()
            return s[len(prefix):].strip() if s.startswith(prefix) else s

        opts = [
            strip_prefix(paragraphs[a_i].text, "A)"),
            strip_prefix(paragraphs[b_i].text, "B)"),
            strip_prefix(paragraphs[c_i].text, "C)"),
            strip_prefix(paragraphs[d_i].text, "D)"),
        ]

        parsed.append({
            "label": label,
            "a_idx": a_i,
            "b_idx": b_i,
            "c_idx": c_i,
            "d_idx": d_i,
            "ans_idx": ans_i,
            "options": opts,
            "answer": answer,
        })

        i = ans_i + 1

    return parsed, unparsed


def build_target_sequence(total, targets, seed=None):
    """
    Build a list of target answer letters of length `total` whose counts
    match `targets` [nA, nB, nC, nD]. The list is shuffled so the pattern
    is not predictable.

    If `seed` is None, uses system randomness (different order each run).
    If `seed` is an int, produces the same order every run.
    """
    letters = ["A", "B", "C", "D"]
    pool = []
    for letter, count in zip(letters, targets):
        pool.extend([letter] * count)

    if len(pool) != total:
        raise RuntimeError(
            f"Target pool length {len(pool)} != question count {total}."
        )

    rng = random.Random(seed) if seed is not None else random
    rng.shuffle(pool)
    return pool


def adjust_options(options, old_answer, new_answer):
    old_idx = ord(old_answer) - ord("A")
    new_idx = ord(new_answer) - ord("A")
    correct = options[old_idx]
    others = [opt for i, opt in enumerate(options) if i != old_idx]
    result = [None] * 4
    result[new_idx] = correct
    it = iter(others)
    for i in range(4):
        if result[i] is None:
            result[i] = next(it)
    return result


def rewrite_paragraph_text(paragraph, new_text):
    """Replace paragraph text while preserving the first run's formatting."""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def insert_unparsed_header(doc, unparsed_labels):
    """Prepend a header block listing the labels of unparsed questions."""
    if not unparsed_labels:
        return
    first = doc.paragraphs[0]
    header = doc.add_paragraph("Questions not parsed:")
    header_el = header._element
    header_el.getparent().remove(header_el)
    first._element.addprevious(header_el)

    anchor = header
    for label in unparsed_labels:
        line = doc.add_paragraph(f"- {label}")
        line_el = line._element
        line_el.getparent().remove(line_el)
        anchor._element.addnext(line_el)
        anchor = line


def process_docx(path):
    doc = Document(str(path))
    paragraphs = doc.paragraphs

    parsed, unparsed = find_questions(paragraphs)

    if not parsed:
        return {"parsed": 0, "unparsed": unparsed, "skipped": True}

    total = len(parsed)
    if sum(TARGETS) != total:
        ratio = TARGETS
        raw = [total * r / sum(ratio) for r in ratio]
        targets = [int(x) for x in raw]
        leftover = total - sum(targets)
        for i in range(leftover):
            targets[i % 4] += 1
    else:
        targets = TARGETS

    sequence = build_target_sequence(total, targets)

    for q, new_letter in zip(parsed, sequence):
        new_options = adjust_options(q["options"], q["answer"], new_letter)
        rewrite_paragraph_text(paragraphs[q["a_idx"]], f"A) {new_options[0]}")
        rewrite_paragraph_text(paragraphs[q["b_idx"]], f"B) {new_options[1]}")
        rewrite_paragraph_text(paragraphs[q["c_idx"]], f"C) {new_options[2]}")
        rewrite_paragraph_text(paragraphs[q["d_idx"]], f"D) {new_options[3]}")
        rewrite_paragraph_text(paragraphs[q["ans_idx"]], f"Answer: {new_letter}")

    insert_unparsed_header(doc, unparsed)

    doc.save(str(path))
    return {"parsed": total, "unparsed": unparsed, "targets": targets}


def main():
    if not BASE_DIR.exists():
        print(f"Base dir not found: {BASE_DIR}")
        sys.exit(1)

    files = sorted(
        p for p in BASE_DIR.iterdir()
        if p.is_file() and p.name.endswith(NAME_FILTER)
    )

    if not files:
        print(f"No files ending in '{NAME_FILTER}' found in {BASE_DIR}")
        sys.exit(1)

    print(f"Found {len(files)} exercise files.")
    for path in files:
        print(f"\n{path.name}")
        try:
            result = process_docx(path)
            if result.get("skipped"):
                print("  skipped (no questions parsed)")
                if result.get("unparsed"):
                    print(f"  unparsed labels: {result['unparsed']}")
                continue
            print(f"  parsed: {result['parsed']}")
            print(
                f"  targets: A={result['targets'][0]} "
                f"B={result['targets'][1]} "
                f"C={result['targets'][2]} "
                f"D={result['targets'][3]}"
            )
            if result["unparsed"]:
                print(f"  unparsed: {len(result['unparsed'])}")
                for u in result["unparsed"]:
                    print(f"    - {u}")
        except Exception as e:
            print(f"  ERROR: {e}")


if __name__ == "__main__":
    main()